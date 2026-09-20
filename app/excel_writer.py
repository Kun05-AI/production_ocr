from __future__ import annotations

from pathlib import Path
from openpyxl import Workbook, load_workbook

SHEET05_COLUMNS = [
    "RunID", "Ngày", "Ca", "Mã đơn hàng", "Mã bản vẽ", "Rev", "Dòng/Lô",
    "Mã dòng hệ thống", "Thứ tự CĐ", "OperationCode", "OperationPlanID", "Nhân viên",
    "Máy", "ActivityType", "Bắt đầu", "Kết thúc", "Nghỉ", "Setup TT", "Downtime",
    "Run TT", "SL xử lý", "SL đạt", "SL đạt lần đầu", "SL cần sửa", "SL phế",
    "SL chờ", "Cân đối SL", "Earned Min", "Labor Min", "Hiệu suất", "Chi phí LĐ",
    "Chi phí máy", "Chi phí chuyển đổi", "Duyệt", "Kiểm tra dữ liệu", "Ghi chú"
]


def write_rows(output_path: str | Path, rows: list[dict], sheet_name: str = "Sheet 05") -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    if output_path.exists():
        wb = load_workbook(output_path)
    else:
        wb = Workbook()
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.create_sheet(sheet_name)
    if ws.max_row == 1 and all(ws.cell(1, i + 1).value is None for i in range(len(SHEET05_COLUMNS))):
        for i, name in enumerate(SHEET05_COLUMNS, 1):
            ws.cell(1, i, name)

    index = {name: i + 1 for i, name in enumerate(SHEET05_COLUMNS)}
    for row in rows:
        excel_row = ws.max_row + 1
        for key, value in row.items():
            if key in index:
                ws.cell(excel_row, index[key], value)
    wb.save(output_path)
    return output_path
